"""Export format tools for Word Document Server."""
import json
import os
import sys
from typing import Optional

from docx import Document

from word_document_server.utils.file_utils import ensure_docx_extension


async def export_to_html(filename: str, output_path: Optional[str] = None) -> str:
    """Export a Word document to HTML. Best results with Microsoft Word (COM). Falls back to basic HTML on other platforms.

    Args:
        filename: Path to the Word document.
        output_path: Output HTML path. Defaults to same name with .html.

    Returns:
        JSON with result.
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})

    if not output_path:
        output_path = os.path.splitext(filename)[0] + ".html"

    if sys.platform == "win32":
        try:
            from word_document_server.core.word_com import comtypes_word_app
            with comtypes_word_app() as (word, _):
                doc = word.Documents.Open(os.path.abspath(filename))
                doc.SaveAs2(os.path.abspath(output_path), FileFormat=8)
                doc.Close()
            return json.dumps({"success": True, "output": output_path, "format": "html"})
        except Exception:
            pass

    try:
        doc = Document(filename)
        html_parts = ["<html><head><meta charset='utf-8'><style>body{font-family:Calibri,sans-serif;margin:40px}table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:4px}</style></head><body>"]
        for para in doc.paragraphs:
            style = para.style.name if para.style else ""
            text = para.text
            if "Heading 1" in style:
                html_parts.append(f"<h1>{text}</h1>")
            elif "Heading 2" in style:
                html_parts.append(f"<h2>{text}</h2>")
            elif "Heading 3" in style:
                html_parts.append(f"<h3>{text}</h3>")
            elif "Title" in style:
                html_parts.append(f"<h1 style='text-align:center'>{text}</h1>")
            else:
                bold = any(run.bold for run in para.runs if run.bold)
                italic = any(run.italic for run in para.runs if run.italic)
                if bold and italic:
                    html_parts.append(f"<p><b><i>{text}</i></b></p>")
                elif bold:
                    html_parts.append(f"<p><b>{text}</b></p>")
                elif italic:
                    html_parts.append(f"<p><i>{text}</i></p>")
                else:
                    html_parts.append(f"<p>{text}</p>")

        for table in doc.tables:
            html_parts.append("<table>")
            for i, row in enumerate(table.rows):
                html_parts.append("<tr>")
                tag = "th" if i == 0 else "td"
                for cell in row.cells:
                    html_parts.append(f"<{tag}>{cell.text}</{tag}>")
                html_parts.append("</tr>")
            html_parts.append("</table>")

        html_parts.append("</body></html>")
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(html_parts))
        return json.dumps({"success": True, "output": output_path, "format": "html", "note": "Basic HTML (no COM). Use live version for full fidelity."})
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to export HTML: {str(e)}"})


async def export_to_rtf(filename: str, output_path: Optional[str] = None) -> str:
    """Export a Word document to RTF. Requires Microsoft Word (COM) on Windows.

    Args:
        filename: Path to the Word document.
        output_path: Output RTF path. Defaults to same name with .rtf.

    Returns:
        JSON with result.
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})

    if not output_path:
        output_path = os.path.splitext(filename)[0] + ".rtf"

    if sys.platform == "win32":
        try:
            from word_document_server.core.word_com import comtypes_word_app
            with comtypes_word_app() as (word, _):
                doc = word.Documents.Open(os.path.abspath(filename))
                doc.SaveAs2(os.path.abspath(output_path), FileFormat=6)
                doc.Close()
            return json.dumps({"success": True, "output": output_path, "format": "rtf"})
        except Exception as e:
            return json.dumps({"success": False, "error": f"Failed to export RTF: {str(e)}"})

    return json.dumps({"success": False, "error": "RTF export requires Microsoft Word on Windows"})


async def export_to_txt(filename: str, output_path: Optional[str] = None) -> str:
    """Export a Word document to plain text.

    Args:
        filename: Path to the Word document.
        output_path: Output text path. Defaults to same name with .txt.

    Returns:
        JSON with result.
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})

    if not output_path:
        output_path = os.path.splitext(filename)[0] + ".txt"

    try:
        doc = Document(filename)
        lines = []
        for para in doc.paragraphs:
            lines.append(para.text)
        for table in doc.tables:
            lines.append("")
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                lines.append("\t".join(cells))
            lines.append("")

        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return json.dumps({"success": True, "output": output_path, "format": "txt"})
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to export text: {str(e)}"})
