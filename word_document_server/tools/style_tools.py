"""Style management tools for Word Document Server."""
import json
import os
from typing import Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor

from word_document_server.utils.file_utils import check_file_writeable, ensure_docx_extension, get_file_lock


def _style_type_name(st):
    mapping = {
        WD_STYLE_TYPE.PARAGRAPH: "paragraph",
        WD_STYLE_TYPE.CHARACTER: "character",
        WD_STYLE_TYPE.TABLE: "table",
        WD_STYLE_TYPE.LIST: "list",
    }
    return mapping.get(st, "unknown")


async def list_styles(filename: str) -> str:
    """List all styles in a Word document with their type and visibility.

    Args:
        filename: Path to the Word document.

    Returns:
        JSON with list of styles.
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})
    try:
        doc = Document(filename)
        styles = []
        for s in doc.styles:
            entry = {
                "name": s.name,
                "type": _style_type_name(s.type),
                "builtin": s.builtin,
                "hidden": s.hidden,
                "priority": s.priority,
            }
            try:
                entry["base_style"] = s.base_style.name if s.base_style else None
            except AttributeError:
                entry["base_style"] = None
            styles.append(entry)
        styles.sort(key=lambda x: (x["builtin"], x["priority"] or 9999, x["name"]))
        return json.dumps({"success": True, "count": len(styles), "styles": styles}, indent=2)
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to list styles: {str(e)}"})


async def get_style_details(filename: str, style_name: str) -> str:
    """Get detailed formatting properties of a named style.

    Args:
        filename: Path to the Word document.
        style_name: Name of the style to inspect.

    Returns:
        JSON with style properties.
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})
    try:
        doc = Document(filename)
        style = doc.styles[style_name]
        info = {
            "name": style.name,
            "type": _style_type_name(style.type),
            "builtin": style.builtin,
            "hidden": style.hidden,
            "priority": style.priority,
        }
        try:
            info["base_style"] = style.base_style.name if style.base_style else None
        except AttributeError:
            info["base_style"] = None
        try:
            info["next_style"] = style.next_style.name if style.next_style else None
        except AttributeError:
            info["next_style"] = None
        f = style.font
        info["font"] = {
            "name": f.name,
            "size_pt": f.size.pt if f.size else None,
            "bold": f.bold,
            "italic": f.italic,
            "underline": f.underline,
            "color_rgb": str(f.color.rgb) if f.color and f.color.rgb else None,
        }
        pf = style.paragraph_format
        align_val = None
        if pf.alignment is not None:
            align_map = {0: "left", 1: "center", 2: "right", 3: "justify"}
            try:
                align_val = align_map.get(int(pf.alignment), str(pf.alignment))
            except (TypeError, ValueError):
                align_val = str(pf.alignment)
        info["paragraph_format"] = {
            "alignment": align_val,
            "space_before_pt": pf.space_before.pt if pf.space_before else None,
            "space_after_pt": pf.space_after.pt if pf.space_after else None,
            "line_spacing": pf.line_spacing,
            "first_line_indent_pt": pf.first_line_indent.pt if pf.first_line_indent else None,
            "keep_with_next": pf.keep_with_next,
            "keep_together": pf.keep_together,
            "page_break_before": pf.page_break_before,
        }
        return json.dumps({"success": True, "style": info}, indent=2)
    except KeyError:
        return json.dumps({"success": False, "error": f"Style '{style_name}' not found"})
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to get style details: {str(e)}"})


async def apply_style(filename: str, paragraph_index: int, style_name: str) -> str:
    """Apply a named style to a specific paragraph.

    Args:
        filename: Path to the Word document.
        paragraph_index: Index of the paragraph (0-based).
        style_name: Name of the style to apply.

    Returns:
        JSON with result.
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return json.dumps({"success": False, "error": f"Cannot modify document: {error_message}"})
    try:
        async with get_file_lock(filename):
            doc = Document(filename)
            if paragraph_index < 0 or paragraph_index >= len(doc.paragraphs):
                return json.dumps({"success": False, "error": f"Paragraph index {paragraph_index} out of range"})
            try:
                style = doc.styles[style_name]
            except KeyError:
                return json.dumps({"success": False, "error": f"Style '{style_name}' not found"})
            para = doc.paragraphs[paragraph_index]
            para.style = style
            doc.save(filename)
            return json.dumps({"success": True, "applied": style_name, "paragraph_index": paragraph_index})
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to apply style: {str(e)}"})


async def modify_style(filename: str, style_name: str, bold: Optional[bool] = None, italic: Optional[bool] = None, font_size: Optional[int] = None, font_name: Optional[str] = None, color: Optional[str] = None, alignment: Optional[str] = None, space_before: Optional[float] = None, space_after: Optional[float] = None, line_spacing: Optional[float] = None) -> str:
    """Modify formatting properties of an existing style.

    Args:
        filename: Path to the Word document.
        style_name: Name of the style to modify.
        bold: Set bold on/off.
        italic: Set italic on/off.
        font_size: Font size in points.
        font_name: Font family name.
        color: Font color — named or hex.
        alignment: Paragraph alignment ("left", "center", "right", "justify").
        space_before: Space before paragraph in points.
        space_after: Space after paragraph in points.
        line_spacing: Line spacing multiplier.

    Returns:
        JSON with result.
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return json.dumps({"success": False, "error": f"Cannot modify document: {error_message}"})
    try:
        async with get_file_lock(filename):
            doc = Document(filename)
            try:
                style = doc.styles[style_name]
            except KeyError:
                return json.dumps({"success": False, "error": f"Style '{style_name}' not found"})

            changes = []
            f = style.font
            pf = style.paragraph_format

            if bold is not None:
                f.bold = bold
                changes.append(f"bold={bold}")
            if italic is not None:
                f.italic = italic
                changes.append(f"italic={italic}")
            if font_size is not None:
                f.size = Pt(font_size)
                changes.append(f"font_size={font_size}pt")
            if font_name is not None:
                f.name = font_name
                changes.append(f"font_name={font_name}")
            if color is not None:
                color_map = {
                    "red": RGBColor(255, 0, 0), "blue": RGBColor(0, 0, 255),
                    "green": RGBColor(0, 128, 0), "black": RGBColor(0, 0, 0),
                    "gray": RGBColor(128, 128, 128), "white": RGBColor(255, 255, 255),
                }
                if color.lower() in color_map:
                    f.color.rgb = color_map[color.lower()]
                else:
                    f.color.rgb = RGBColor.from_string(color)
                changes.append(f"color={color}")

            align_map = {"left": 0, "center": 1, "right": 2, "justify": 3}
            if alignment is not None:
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                pf.alignment = WD_ALIGN_PARAGRAPH(align_map.get(alignment.lower(), 0))
                changes.append(f"alignment={alignment}")
            if space_before is not None:
                pf.space_before = Pt(space_before)
                changes.append(f"space_before={space_before}pt")
            if space_after is not None:
                pf.space_after = Pt(space_after)
                changes.append(f"space_after={space_after}pt")
            if line_spacing is not None:
                pf.line_spacing = line_spacing
                changes.append(f"line_spacing={line_spacing}")

            if not changes:
                return json.dumps({"success": False, "error": "No changes specified"})

            doc.save(filename)
            return json.dumps({"success": True, "style": style_name, "changes": changes})
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to modify style: {str(e)}"})


async def delete_style(filename: str, style_name: str) -> str:
    """Delete a custom (non-built-in) style from a document.

    Args:
        filename: Path to the Word document.
        style_name: Name of the style to delete.

    Returns:
        JSON with result.
    """
    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return json.dumps({"success": False, "error": f"Cannot modify document: {error_message}"})
    try:
        async with get_file_lock(filename):
            doc = Document(filename)
            try:
                style = doc.styles[style_name]
            except KeyError:
                return json.dumps({"success": False, "error": f"Style '{style_name}' not found"})
            if style.builtin:
                return json.dumps({"success": False, "error": f"Cannot delete built-in style '{style_name}'"})

            styles_element = doc.styles.element
            style_element = style.element
            styles_element.remove(style_element)
            doc.save(filename)
            return json.dumps({"success": True, "deleted": style_name})
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to delete style: {str(e)}"})


async def copy_styles_from_template(filename: str, template_path: str) -> str:
    """Copy styles from another document or template into the current document.

    Args:
        filename: Path to the target Word document.
        template_path: Path to the source document/template to copy styles from.

    Returns:
        JSON with result.
    """
    filename = ensure_docx_extension(filename)
    template_path = ensure_docx_extension(template_path)
    if not os.path.exists(filename):
        return json.dumps({"success": False, "error": f"Document {filename} does not exist"})
    if not os.path.exists(template_path):
        return json.dumps({"success": False, "error": f"Template {template_path} does not exist"})
    is_writeable, error_message = check_file_writeable(filename)
    if not is_writeable:
        return json.dumps({"success": False, "error": f"Cannot modify document: {error_message}"})
    try:
        async with get_file_lock(filename):
            doc = Document(filename)
            tmpl = Document(template_path)

            imported = []
            skipped = []
            for s in tmpl.styles:
                if s.builtin:
                    continue
                try:
                    doc.styles[s.name]
                    skipped.append(s.name)
                    continue
                except KeyError:
                    pass
                try:
                    new_s = doc.styles.add_style(s.name, s.type)
                    if s.base_style:
                        try:
                            new_s.base_style = doc.styles[s.base_style.name]
                        except KeyError:
                            pass
                    new_s.font.name = s.font.name
                    if s.font.size:
                        new_s.font.size = s.font.size
                    new_s.font.bold = s.font.bold
                    new_s.font.italic = s.font.italic
                    if s.font.color and s.font.color.rgb:
                        new_s.font.color.rgb = s.font.color.rgb
                    if s.paragraph_format.alignment is not None:
                        new_s.paragraph_format.alignment = s.paragraph_format.alignment
                    if s.paragraph_format.space_before:
                        new_s.paragraph_format.space_before = s.paragraph_format.space_before
                    if s.paragraph_format.space_after:
                        new_s.paragraph_format.space_after = s.paragraph_format.space_after
                    imported.append(s.name)
                except Exception:
                    skipped.append(s.name)

            doc.save(filename)
            return json.dumps({
                "success": True,
                "imported": imported,
                "imported_count": len(imported),
                "skipped": skipped,
                "skipped_count": len(skipped),
            })
    except Exception as e:
        return json.dumps({"success": False, "error": f"Failed to copy styles: {str(e)}"})
