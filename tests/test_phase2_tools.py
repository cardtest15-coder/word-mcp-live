"""Tests for Phase 2 file-based tools: export, merge, highlight, field, property, page_design."""
import asyncio
import json
import os
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches, Pt

from word_document_server.tools.export_tools import export_to_txt
from word_document_server.tools.merge_tools import mail_merge
from word_document_server.tools.highlight_tools import highlight_text, remove_highlight
from word_document_server.tools.field_tools import insert_field, insert_content_control
from word_document_server.tools.property_tools import get_custom_properties, set_custom_property
from word_document_server.tools.page_design_tools import (
    set_different_first_page, set_odd_even_headers, set_tab_stops,
    clear_tab_stops, insert_drop_cap, set_page_borders,
)
from word_document_server.tools.style_tools import list_styles, get_style_details, apply_style, modify_style
from word_document_server.tools.column_tools import get_section_layout, set_section_columns, insert_column_break
from word_document_server.tools.advanced_table_tools import (
    delete_table, repeat_table_header, convert_table_to_text,
    convert_text_to_table, sort_table,
)


def _make_sample_docx(path: Path, content: str = "Hello world") -> None:
    doc = Document()
    doc.add_heading("Test Document", level=1)
    doc.add_paragraph(content)
    doc.save(path)


def _make_merge_template(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Dear <<Name>>, your order <<Order>> is ready.")
    doc.save(path)


def _run(coro):
    return asyncio.run(coro)


# ── Export Tools ──────────────────────────────────────────────────────

class TestExportToTxt:
    def test_export_to_txt(self, tmp_path: Path):
        src = tmp_path / "test.docx"
        _make_sample_docx(src, "Export me")
        out = tmp_path / "test.txt"
        result = json.loads(_run(export_to_txt(str(src), str(out))))
        assert result["success"]
        assert out.exists()
        assert "Export me" in out.read_text(encoding="utf-8")

    def test_export_to_txt_missing_file(self):
        result = json.loads(_run(export_to_txt("nonexistent.docx")))
        assert not result["success"]

    def test_export_to_txt_with_table(self, tmp_path: Path):
        src = tmp_path / "table.docx"
        doc = Document()
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "A"
        table.rows[0].cells[1].text = "B"
        table.rows[1].cells[0].text = "1"
        table.rows[1].cells[1].text = "2"
        doc.save(src)
        out = tmp_path / "table.txt"
        result = json.loads(_run(export_to_txt(str(src), str(out))))
        assert result["success"]
        content = out.read_text(encoding="utf-8")
        assert "Before table" in content
        assert "A\tB" in content


# ── Mail Merge ────────────────────────────────────────────────────────

class TestMailMerge:
    def test_mail_merge_basic(self, tmp_path: Path):
        template = tmp_path / "template.docx"
        _make_merge_template(template)
        data = "Name,Order\nAlice,123\nBob,456"
        result = json.loads(_run(mail_merge(str(template), data)))
        assert result["success"]
        assert result["rows"] == 2
        merged = tmp_path / "template_merged.docx"
        assert merged.exists()
        doc = Document(str(merged))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Alice" in text
        assert "123" in text

    def test_mail_merge_custom_output(self, tmp_path: Path):
        template = tmp_path / "tmpl.docx"
        _make_merge_template(template)
        out = tmp_path / "result.docx"
        result = json.loads(_run(mail_merge(str(template), "Name,Order\nEve,789", str(out))))
        assert result["success"]
        assert out.exists()

    def test_mail_merge_no_fields(self, tmp_path: Path):
        template = tmp_path / "plain.docx"
        _make_sample_docx(template, "No merge fields here")
        result = json.loads(_run(mail_merge(str(template), "Name\nAlice")))
        assert not result["success"]

    def test_mail_merge_insufficient_data(self, tmp_path: Path):
        template = tmp_path / "tmpl.docx"
        _make_merge_template(template)
        result = json.loads(_run(mail_merge(str(template), "Name,Order")))
        assert not result["success"]

    def test_mail_merge_separator(self, tmp_path: Path):
        template = tmp_path / "tmpl.docx"
        _make_merge_template(template)
        data = "Name;Order\nCharlie;999"
        result = json.loads(_run(mail_merge(str(template), data, separator="semicolon")))
        assert result["success"]


# ── Highlight Tools ──────────────────────────────────────────────────

class TestHighlightTools:
    def test_highlight_text(self, tmp_path: Path):
        src = tmp_path / "hl.docx"
        _make_sample_docx(src, "Highlight this text please")
        result = json.loads(_run(highlight_text(str(src), 1, 0, 9, "yellow")))
        assert result["success"]
        assert result["color"] == "yellow"

    def test_highlight_text_bad_color(self, tmp_path: Path):
        src = tmp_path / "hl.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(highlight_text(str(src), 1, 0, 4, "neon")))
        assert not result["success"]
        assert "Unknown color" in result["error"]

    def test_highlight_text_bad_range(self, tmp_path: Path):
        src = tmp_path / "hl.docx"
        _make_sample_docx(src, "Short")
        result = json.loads(_run(highlight_text(str(src), 1, 0, 999, "yellow")))
        assert not result["success"]

    def test_remove_highlight(self, tmp_path: Path):
        src = tmp_path / "hl.docx"
        _make_sample_docx(src, "Remove my highlight")
        _run(highlight_text(str(src), 1, 0, 6, "yellow"))
        result = json.loads(_run(remove_highlight(str(src))))
        assert result["success"]
        assert result["removed_count"] >= 1

    def test_remove_highlight_single_para(self, tmp_path: Path):
        src = tmp_path / "hl.docx"
        doc = Document()
        doc.add_paragraph("First para highlight me")
        doc.add_paragraph("Second para also highlight")
        doc.save(src)
        _run(highlight_text(str(src), 0, 0, 5, "green"))
        _run(highlight_text(str(src), 1, 0, 6, "red"))
        result = json.loads(_run(remove_highlight(str(src), paragraph_index=0)))
        assert result["success"]
        assert result["scope"] == "paragraph 0"


# ── Field Tools ────────────────────────────────────────────────────────

class TestFieldTools:
    def test_insert_field(self, tmp_path: Path):
        src = tmp_path / "field.docx"
        _make_sample_docx(src, "Page numbers here")
        result = json.loads(_run(insert_field(str(src), "PAGE")))
        assert result["success"]
        assert result["field"] == "PAGE"

    def test_insert_field_invalid_type(self, tmp_path: Path):
        src = tmp_path / "field.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(insert_field(str(src), "INVALID")))
        assert not result["success"]
        assert "Unknown field" in result["error"]

    def test_insert_field_with_display(self, tmp_path: Path):
        src = tmp_path / "field.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(insert_field(str(src), "DATE", field_text="2025-01-01")))
        assert result["success"]

    def test_insert_content_control(self, tmp_path: Path):
        src = tmp_path / "cc.docx"
        _make_sample_docx(src, "Fill this in")
        result = json.loads(_run(insert_content_control(str(src), "text", title="Name", placeholder="Enter name")))
        assert result["success"]
        assert result["control_type"] == "text"

    def test_insert_content_control_dropdown(self, tmp_path: Path):
        src = tmp_path / "cc.docx"
        _make_sample_docx(src, "Choose")
        result = json.loads(_run(insert_content_control(str(src), "dropdown", title="Select")))
        assert result["success"]

    def test_insert_content_control_invalid(self, tmp_path: Path):
        src = tmp_path / "cc.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(insert_content_control(str(src), "spinner")))
        assert not result["success"]


# ── Property Tools ────────────────────────────────────────────────────

class TestPropertyTools:
    def test_get_custom_properties(self, tmp_path: Path):
        src = tmp_path / "props.docx"
        doc = Document()
        doc.core_properties.author = "Test Author"
        doc.core_properties.title = "Test Title"
        doc.save(src)
        result = json.loads(_run(get_custom_properties(str(src))))
        assert result["success"]
        assert result["properties"]["author"] == "Test Author"
        assert result["properties"]["title"] == "Test Title"

    def test_set_core_property(self, tmp_path: Path):
        src = tmp_path / "props.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(set_custom_property(str(src), "title", "My Doc")))
        assert result["success"]
        assert result["type"] == "core"
        doc = Document(str(src))
        assert doc.core_properties.title == "My Doc"

    def test_set_custom_property_text(self, tmp_path: Path):
        src = tmp_path / "props.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(set_custom_property(str(src), "Project", "Alpha", "text")))
        assert result["success"]
        assert result["property"] == "Project"

    def test_set_custom_property_auto_type(self, tmp_path: Path):
        src = tmp_path / "props.docx"
        _make_sample_docx(src, "Test")
        result_int = json.loads(_run(set_custom_property(str(src), "Version", 3)))
        assert result_int["success"]
        assert result_int["type"] == "number"
        result_bool = json.loads(_run(set_custom_property(str(src), "Reviewed", True)))
        assert result_bool["success"]
        assert result_bool["type"] == "boolean"


# ── Page Design Tools ──────────────────────────────────────────────────

class TestPageDesignTools:
    def test_set_different_first_page(self, tmp_path: Path):
        src = tmp_path / "design.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(set_different_first_page(str(src), 0, True)))
        assert result["success"]
        assert result["different_first_page"] is True

    def test_set_different_first_page_bad_section(self, tmp_path: Path):
        src = tmp_path / "design.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(set_different_first_page(str(src), 99, True)))
        assert not result["success"]

    def test_set_odd_even_headers(self, tmp_path: Path):
        src = tmp_path / "design.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(set_odd_even_headers(str(src), 0, True)))
        assert result["success"]

    def test_set_tab_stops(self, tmp_path: Path):
        src = tmp_path / "tabs.docx"
        _make_sample_docx(src, "Left\tCenter\tRight")
        result = json.loads(_run(set_tab_stops(str(src), 1, "0.5,2.0,4.0", "left,center,right", "none,dots,none")))
        assert result["success"]
        assert result["tabs_added"] == 3

    def test_set_tab_stops_default_align(self, tmp_path: Path):
        src = tmp_path / "tabs.docx"
        _make_sample_docx(src, "Tab test")
        result = json.loads(_run(set_tab_stops(str(src), 1, "1.0,3.0")))
        assert result["success"]
        assert result["tabs_added"] == 2

    def test_clear_tab_stops(self, tmp_path: Path):
        src = tmp_path / "tabs.docx"
        _make_sample_docx(src, "Tab test")
        _run(set_tab_stops(str(src), 1, "1.0,3.0"))
        result = json.loads(_run(clear_tab_stops(str(src), 1)))
        assert result["success"]
        assert result["tabs_cleared"] is True

    def test_insert_drop_cap(self, tmp_path: Path):
        src = tmp_path / "dropcap.docx"
        _make_sample_docx(src, "Once upon a time")
        result = json.loads(_run(insert_drop_cap(str(src), 1, lines=3)))
        assert result["success"]
        assert result["drop_char"] == "O"

    def test_insert_drop_cap_empty_para(self, tmp_path: Path):
        src = tmp_path / "dropcap.docx"
        doc = Document()
        doc.add_paragraph("")
        doc.save(src)
        result = json.loads(_run(insert_drop_cap(str(src), 0)))
        assert not result["success"]

    def test_set_page_borders(self, tmp_path: Path):
        src = tmp_path / "borders.docx"
        _make_sample_docx(src, "Bordered page")
        result = json.loads(_run(set_page_borders(str(src), 0, "single", "auto", 4, 24.0)))
        assert result["success"]

    def test_set_page_borders_remove(self, tmp_path: Path):
        src = tmp_path / "borders.docx"
        _make_sample_docx(src, "Test")
        _run(set_page_borders(str(src), 0, "single"))
        result = json.loads(_run(set_page_borders(str(src), 0, "none")))
        assert result["success"]
        assert result["borders"] == "removed"

    def test_set_page_borders_bad_style(self, tmp_path: Path):
        src = tmp_path / "borders.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(set_page_borders(str(src), 0, "wavy")))
        assert not result["success"]


# ── Style Tools (Phase 1) ─────────────────────────────────────────────

class TestStyleTools:
    def test_list_styles(self, tmp_path: Path):
        src = tmp_path / "styles.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(list_styles(str(src))))
        assert result["success"]
        assert result["count"] > 0

    def test_get_style_details(self, tmp_path: Path):
        src = tmp_path / "styles.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(get_style_details(str(src), "Normal")))
        assert result["success"]

    def test_apply_style(self, tmp_path: Path):
        src = tmp_path / "styles.docx"
        doc = Document()
        doc.add_paragraph("Style me")
        doc.save(src)
        result = json.loads(_run(apply_style(str(src), 0, "Heading 1")))
        assert result["success"]

    def test_modify_style(self, tmp_path: Path):
        src = tmp_path / "styles.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(modify_style(str(src), "Normal", font_size=14)))
        assert result["success"]


# ── Column Tools (Phase 1) ─────────────────────────────────────────────

class TestColumnTools:
    def test_get_section_layout(self, tmp_path: Path):
        src = tmp_path / "cols.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(get_section_layout(str(src))))
        assert result["success"]

    def test_set_section_columns(self, tmp_path: Path):
        src = tmp_path / "cols.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(set_section_columns(str(src), 0, columns=2)))
        assert result["success"]
        assert result["columns"] == 2

    def test_insert_column_break(self, tmp_path: Path):
        src = tmp_path / "cols.docx"
        _make_sample_docx(src, "Test")
        result = json.loads(_run(insert_column_break(str(src))))
        assert result["success"]


# ── Advanced Table Tools (Phase 1) ─────────────────────────────────────

class TestAdvancedTableTools:
    def test_delete_table(self, tmp_path: Path):
        src = tmp_path / "table.docx"
        doc = Document()
        doc.add_paragraph("Before")
        doc.add_table(rows=2, cols=2)
        doc.add_paragraph("After")
        doc.save(src)
        result = json.loads(_run(delete_table(str(src), 0)))
        assert result["success"]

    def test_repeat_table_header(self, tmp_path: Path):
        src = tmp_path / "table.docx"
        doc = Document()
        doc.add_table(rows=3, cols=2)
        doc.save(src)
        result = json.loads(_run(repeat_table_header(str(src), 0, True)))
        assert result["success"]

    def test_convert_table_to_text(self, tmp_path: Path):
        src = tmp_path / "table.docx"
        doc = Document()
        t = doc.add_table(rows=2, cols=2)
        t.rows[0].cells[0].text = "A"
        t.rows[0].cells[1].text = "B"
        t.rows[1].cells[0].text = "1"
        t.rows[1].cells[1].text = "2"
        doc.save(src)
        result = json.loads(_run(convert_table_to_text(str(src), 0)))
        assert result["success"]
        assert "A" in result["text"]

    def test_convert_text_to_table(self, tmp_path: Path):
        src = tmp_path / "table.docx"
        doc = Document()
        doc.add_paragraph("Keep this")
        doc.save(src)
        result = json.loads(_run(convert_text_to_table(str(src), "A,B\n1,2", 2, ",")))
        assert result["success"]

    def test_sort_table(self, tmp_path: Path):
        src = tmp_path / "table.docx"
        doc = Document()
        t = doc.add_table(rows=3, cols=1)
        t.rows[0].cells[0].text = "Charlie"
        t.rows[1].cells[0].text = "Alice"
        t.rows[2].cells[0].text = "Bob"
        doc.save(src)
        from unittest.mock import patch, MagicMock
        mock_app = MagicMock()
        mock_doc = MagicMock()
        mock_tbl = MagicMock()
        mock_doc.Tables.Count = 1
        mock_doc.Tables.return_value = mock_tbl
        with patch("word_document_server.core.word_com.get_word_app", return_value=mock_app), \
             patch("word_document_server.core.word_com.find_document", return_value=mock_doc):
            result = json.loads(_run(sort_table(str(src), 0, 0)))
        assert result["success"]
        mock_tbl.Sort.assert_called_once()
